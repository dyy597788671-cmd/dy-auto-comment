import copy
import json
from pathlib import Path
import struct
import sys
import tempfile
import unittest
from unittest.mock import Mock, patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import core_runner as core
import lan_bridge
import mobile_instances as instances
import module1
from openpyxl import load_workbook
from test_execution_output import temporary_root, prepare_root, finish_kwargs


PACKAGE = instances.APP_PACKAGE


class Shell:
    def __init__(self, users=(0, 735), installed=None):
        self.users = users
        self.installed = set(users if installed is None else installed)
        self.calls = []
        self.foreground = self.users[0]

    def __call__(self, phone, args):
        self.calls.append((phone, args))
        if args == ["pm", "list", "users"]:
            return "Users:\n" + "\n".join("UserInfo{%d:空间%d:%s} running" % (u, u, "c13" if i == 0 else "10")
                                         for i, u in enumerate(self.users))
        if args[:3] == ["pm", "list", "packages"]:
            return "package:" + PACKAGE if args[args.index("--user") + 1] in self.installed else ""
        if args[:3] == ["cmd", "package", "resolve-activity"]:
            return PACKAGE + "/.MainActivity\n"
        if args[:2] == ["am", "start"]:
            self.foreground = args[args.index("--user") + 1]
            return "Starting: Intent {...}\n"
        if args[:2] == ["dumpsys", "activity"]:
            return "topResumedActivity=ActivityRecord{a12 u%d %s/.MainActivity t123}" % (self.foreground, PACKAGE)
        raise AssertionError(args)


def roster(root, shell=None, count=1):
    return instances.discover_roster([{"udid": "usb-%d" % i, "custom_name": "手机%d" % i} for i in range(count)],
                                     root, shell=shell or Shell())


class DiscoveryTests(unittest.TestCase):
    def test_four_phones_seven_instances_keep_shared_connection_indices(self):
        with tempfile.TemporaryDirectory() as root:
            shell = Shell()
            def run(phone, args):
                if phone == "usb-3" and args == ["pm", "list", "users"]:
                    return "UserInfo{0:主空间:c13} running"
                return shell(phone, args)
            values = roster(root, run, 4)
            self.assertEqual(len(values), 7)
            self.assertEqual([v["connection_index"] for v in values], [0, 0, 1, 1, 2, 2, 3])
            self.assertEqual(len({v["sender_account_id"] for v in values}), 7)
            self.assertEqual(values[0]["sender_account_id"], "设备_usb-0")
            self.assertEqual(values[1]["android_user_id"], 735)
            self.assertFalse(any(args[:2] == ["am", "start"] for _, args in shell.calls))

    def test_main_user_identity_survives_order_and_current_user_changes(self):
        with tempfile.TemporaryDirectory() as root:
            before = roster(root, Shell((10, 735)))
            after = roster(root, Shell((735, 10)))
            self.assertEqual(before[0]["sender_account_id"], after[0]["sender_account_id"])
            self.assertEqual(after[0]["android_user_id"], 10)

    def test_users_without_douyin_are_not_counted(self):
        with tempfile.TemporaryDirectory() as root:
            self.assertEqual(len(roster(root, Shell(installed=(0,)))), 1)

    def test_query_failure_removes_stale_available_instances(self):
        with tempfile.TemporaryDirectory() as root:
            roster(root)
            self.assertEqual(roster(root, Mock(side_effect=RuntimeError("空间权限不足"))), [])
            inventory = instances.cached_inventory(root, {"usb-0"})
            self.assertEqual(inventory["instance_count"], 0)
            self.assertIn("空间权限不足", inventory["discovery_errors"][0])

    def test_offline_phone_instances_are_not_counted(self):
        with tempfile.TemporaryDirectory() as root:
            roster(root, count=2)
            inventory = instances.cached_inventory(root, {"usb-1"})
            self.assertEqual(inventory["instance_count"], 2)
            self.assertTrue(all(v["device_id"] == "usb-1" for v in inventory["instances"]))

    def test_web_counts_seven_instances_without_executing_phone_commands(self):
        with tempfile.TemporaryDirectory() as root:
            shell = Shell()
            def run(phone, args):
                if phone == "usb-3" and args == ["pm", "list", "users"]:
                    return "UserInfo{0:主空间:c13} running"
                return shell(phone, args)
            roster(root, run, 4)
            with patch.object(lan_bridge, "_online_registered_devices", return_value={"usb-%d" % i for i in range(4)}), \
                    patch.object(instances, "AdbShell", side_effect=AssertionError("网页不能执行手机命令")):
                result = lan_bridge.connected_inventory(root)
            self.assertEqual(result["connected_device_count"], 7)
            self.assertEqual(result["connected_physical_device_count"], 4)

    def test_missing_inventory_does_not_guess_one_or_two_apps_per_phone(self):
        with tempfile.TemporaryDirectory() as root, \
                patch.object(lan_bridge, "_online_registered_devices", return_value={"a", "b"}):
            result = lan_bridge.connected_inventory(root)
            self.assertIsNone(result["connected_device_count"])
            self.assertEqual(result["connected_physical_device_count"], 2)

    def test_registry_updates_do_not_modify_existing_business_files(self):
        with tempfile.TemporaryDirectory() as root:
            data = Path(root) / "data"
            data.mkdir()
            (data / "runtime_state.json").write_bytes(b"original-state")
            (data / "task_cases.xlsx").write_bytes(b"original-tasks")
            roster(root)
            self.assertEqual((data / "runtime_state.json").read_bytes(), b"original-state")
            self.assertEqual((data / "task_cases.xlsx").read_bytes(), b"original-tasks")


class LaunchTests(unittest.TestCase):
    def test_selects_requested_space_and_confirms_it(self):
        with tempfile.TemporaryDirectory() as root:
            shell = Shell()
            clone = roster(root, shell)[1]
            device = Mock()
            device.get_udid.return_value = "usb-0"
            result = instances.launch_instance(device, clone, shell=shell)
            self.assertEqual(result["android_user_id"], 735)
            start = next(args for _, args in shell.calls if args[:2] == ["am", "start"])
            self.assertEqual(start[start.index("--user") + 1], 735)
            device.open_app.assert_not_called()

    def test_wrong_front_user_stops_before_mobile_actions(self):
        with tempfile.TemporaryDirectory() as root:
            clone = roster(root)[1]
            device = Mock()
            device.get_udid.return_value = "usb-0"
            def shell(phone, args):
                if args[:3] == ["cmd", "package", "resolve-activity"]:
                    return PACKAGE + "/.MainActivity"
                if args[:2] == ["am", "start"]:
                    return "Starting: Intent"
                return "topResumedActivity=ActivityRecord{a u0 " + PACKAGE + "/.MainActivity t1}"
            with self.assertRaisesRegex(RuntimeError, "前台实例核对失败"):
                instances.launch_instance(device, clone, shell=shell, sleep=lambda _: None,
                                          monotonic=Mock(side_effect=[0, 1, 9]))
            device.find_by_xpath.assert_not_called()

    def test_wrong_physical_phone_never_launches(self):
        device = Mock()
        device.get_udid.return_value = "other-phone"
        shell = Mock()
        with self.assertRaisesRegex(RuntimeError, "物理手机不一致"):
            instances.launch_instance(device, {"device_id": "usb-0"}, shell=shell)
        shell.assert_not_called()

    def test_history_activity_is_not_treated_as_foreground(self):
        self.assertIsNone(instances.foreground("Hist #1: ActivityRecord{a u735 " + PACKAGE + "/.MainActivity}"))

    def test_foreground_requires_package_and_user(self):
        result = instances.foreground("mResumedActivity: ActivityRecord{a u10 com.android.settings/.Settings t1}")
        self.assertEqual(result, (10, "com.android.settings"))


class ProtocolTests(unittest.TestCase):
    def test_reads_split_shell_v2_frames_and_propagates_exit_error(self):
        output = "空间735".encode()
        packets = b"OKAYOKAY" + struct.pack("<BI", 1, len(output)) + output + struct.pack("<BI", 3, 1) + b"\x00"
        connection = Mock()
        chunks = bytearray(packets)
        def receive(count):
            value = bytes(chunks[:min(count, 2)])
            del chunks[:len(value)]
            return value
        connection.recv.side_effect = receive
        connection.__enter__ = Mock(return_value=connection)
        connection.__exit__ = Mock(return_value=False)
        shell = object.__new__(instances.AdbShell)
        shell.online = {"usb"}
        with patch.object(instances.socket, "create_connection", return_value=connection):
            self.assertEqual(shell("usb", ["pm", "list", "users"]), "空间735")
        self.assertEqual(connection.sendall.call_args_list[0].args[0], b"0012host:transport:usb")
        error = b"Permission denied"
        chunks = bytearray(b"OKAYOKAY" + struct.pack("<BI", 2, len(error)) + error + struct.pack("<BI", 3, 1) + b"\x01")
        with patch.object(instances.socket, "create_connection", return_value=connection):
            with self.assertRaisesRegex(RuntimeError, "Permission denied"):
                shell("usb", ["pm", "list", "users"])


class TaskTests(unittest.TestCase):
    def prepare(self, root):
        values = roster(root)
        prepare_root(root, []) if not (root / "config").exists() else None
        path = root / "data" / "task_cases.xlsx"
        wb = load_workbook(path)
        wb.active.cell(1, 11, "Target_Device_Count")
        wb.active.cell(2, 11, 2)
        wb.active.cell(2, 10, "评论一\n评论二\n评论三")
        wb.save(path)
        wb.close()
        return values

    def test_one_phone_two_instances_allocate_and_resume_independently(self):
        with temporary_root() as folder:
            root = Path(folder)
            prepare_root(root, [])
            values = self.prepare(root)
            primary, clone = values
            gateway = core.UnifiedGateway(folder, primary["sender_account_id"])
            progress = gateway.allocate_task("1005", values)
            self.assertEqual(progress["assigned_device_count"], 2)
            first = gateway.get_next_task(physical_device_id=primary["device_id"])
            plan = gateway.generate_strategy(seed=17, task_id=first["task_id"], run_token=first["run_token"])
            saved = [{"action": plan["execution_order"][0], "status": "PASSED", "message": ""}]
            gateway.save_action_results(first["task_id"], first["run_token"], saved)
            second_gateway = core.UnifiedGateway(folder, clone["sender_account_id"])
            with self.assertRaisesRegex(core.EngineError, "物理设备"):
                second_gateway.get_next_task(physical_device_id=clone["device_id"])
            gateway.record_result(**finish_kwargs(first, saved, "FAILED"))
            second = second_gateway.get_next_task(physical_device_id=clone["device_id"])
            self.assertEqual(second["android_user_id"], 735)
            self.assertNotEqual(first["comment_content"], second["comment_content"])
            second_plan = second_gateway.generate_strategy(seed=18, task_id=second["task_id"], run_token=second["run_token"])
            done = [{"action": v, "status": "PASSED", "message": ""} for v in second_plan["execution_order"]]
            second_gateway.save_action_results(second["task_id"], second["run_token"], done)
            second_gateway.record_result(**finish_kwargs(second, done))
            gateway.request_resume("1005")
            resumed = gateway.get_next_task(physical_device_id=primary["device_id"])
            self.assertEqual(resumed["action_results"], saved)
            self.assertEqual(resumed["comment_content"], first["comment_content"])
            self.assertEqual(gateway.remaining_strategy(resumed["task_id"], resumed["run_token"])["execution_order"],
                             plan["execution_order"][1:])
            statuses = {v["sender_account_id"]: v["status"] for v in gateway.get_project_progress("1005")["accounts"]}
            self.assertEqual(statuses[clone["sender_account_id"]], "SUCCESS")

    def test_duplicate_instance_with_different_accounts_is_rejected(self):
        with tempfile.TemporaryDirectory() as root:
            values = roster(root)
            duplicate = dict(values[0], sender_account_id="another")
            with self.assertRaises(core.EngineError):
                core._connected_devices([values[0], duplicate])

    def test_clone_is_automatically_available_with_old_sender_account_setting(self):
        with temporary_root() as folder:
            root = Path(folder)
            prepare_root(root, ["设备_usb-0"])
            values = self.prepare(root)
            gateway = core.UnifiedGateway(folder, values[0]["sender_account_id"])
            self.assertEqual(gateway.allocate_task("1005", values)["assigned_device_count"], 2)
            clone = core.UnifiedGateway(folder, values[1]["sender_account_id"])
            self.assertEqual(clone.get_next_task(physical_device_id="usb-0")["android_user_id"], 735)

    def test_main_cooldown_does_not_block_clone_on_same_phone(self):
        with temporary_root() as folder:
            root = Path(folder)
            prepare_root(root, [])
            values = self.prepare(root)
            settings = root / "config/settings.json"
            data = json.loads(settings.read_text())
            data.update(cooldown_min_seconds=300, cooldown_max_seconds=300)
            settings.write_text(json.dumps(data))
            gateway = core.UnifiedGateway(folder, values[0]["sender_account_id"])
            gateway.allocate_task("1005", values)
            main = gateway.get_next_task(physical_device_id="usb-0")
            gateway.record_result(**finish_kwargs(main, [], "NOT_FOUND"))
            selection = gateway.select_dispatch(values)
            self.assertEqual(selection["accounts"], [values[1]["sender_account_id"]])
            self.assertEqual(selection["wait_seconds"], 0)

    def test_old_assignments_keep_main_account_when_clone_is_added(self):
        with temporary_root() as folder:
            root = Path(folder)
            prepare_root(root, [])
            values = self.prepare(root)
            path = root / "data" / "task_cases.xlsx"
            wb = load_workbook(path)
            wb.active.cell(2, 11, 1)
            wb.save(path)
            wb.close()
            gateway = core.UnifiedGateway(folder, "设备_usb-0")
            old = gateway.allocate_task("1005", [{"device_id": "usb-0", "sender_account_id": "设备_usb-0"}])
            self.assertEqual(gateway.allocate_task("1005", values), old)

    def test_main_flow_maps_both_instances_to_original_phone_object(self):
        with temporary_root() as folder:
            root = Path(folder)
            prepare_root(root, [])
            values = self.prepare(root)
            with patch.object(module1, "PROJECT_ROOT", folder), patch.object(module1, "_device_bindings", return_value=values):
                selected = module1.prepare_devices([{"udid": "usb-0"}])
            self.assertEqual([v["connection_index"] for v in selected], [0, 0])


if __name__ == "__main__":
    unittest.main()
